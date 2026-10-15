import io
import os
import json
import csv
import threading
import zipfile
import xml.etree.ElementTree as ET
from flask import Flask, request, jsonify, render_template, send_file
import simplekml

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

app = Flask(__name__)
BASE_DIR = os.path.dirname(__file__)
UPLOAD_DIR = os.path.join(BASE_DIR, 'uploads')
DATA_DIR = os.path.join(BASE_DIR, 'data')
STATE_FILE = os.path.join(DATA_DIR, 'state.json')
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)

# ============================ STATE ============================
ROUTES = {}
_next_id = 1
_state_lock = threading.RLock()          # RLock: позволяет вложенный захват
PERSIST = {'customFolders': set()}

# GPON-состояние: geo-словарь, список ONU и имена загруженных файлов.
# Живёт в памяти и персистится в data/state.json рядом с KMZ.
GPON = {'geo': {}, 'onus': [], 'geoFiles': [], 'vgFile': None}


def _save_state():
    # ВНИМАНИЕ: _save_state() сама захватывает _state_lock (RLock → реентерабельно).
    with _state_lock:
        payload = {
            'next_id': _next_id,
            'customFolders': sorted(PERSIST['customFolders']),
            'routes': list(ROUTES.values()),
            'gpon': {
                'geo': dict(GPON['geo']),
                'onus': list(GPON['onus']),
                'geoFiles': list(GPON.get('geoFiles', [])),
                'vgFile': GPON.get('vgFile'),
            },
        }
        tmp = STATE_FILE + '.tmp'
        with open(tmp, 'w', encoding='utf-8') as f:
            json.dump(payload, f, ensure_ascii=False)
        os.replace(tmp, STATE_FILE)


def _load_state():
    global _next_id

    # Сброс GPON в дефолт — на случай, если файла нет или он битый.
    GPON['geo'] = {}
    GPON['onus'] = []
    GPON['geoFiles'] = []
    GPON['vgFile'] = None

    if not os.path.exists(STATE_FILE):
        return
    try:
        with open(STATE_FILE, 'r', encoding='utf-8') as f:
            payload = json.load(f)
    except (json.JSONDecodeError, OSError) as e:
        print('state load error:', e)
        return

    ROUTES.clear()
    for r in payload.get('routes', []):
        try:
            rid = int(r['id'])
        except (KeyError, TypeError, ValueError):
            continue
        ROUTES[rid] = {
            'id': rid,
            'name': r.get('name', ''),
            'branch': r.get('branch'),
            'folder': r.get('folder', ''),
            'visible': bool(r.get('visible', True)),
            'points': r.get('points', []),
        }

    PERSIST['customFolders'] = set(payload.get('customFolders', []))
    _next_id = int(payload.get('next_id', (max(ROUTES) + 1) if ROUTES else 1))

    # Восстанавливаем GPON (если секция есть).
    gpon = payload.get('gpon') or {}
    GPON['geo'] = gpon.get('geo') or {}
    GPON['onus'] = gpon.get('onus') or []
    GPON['geoFiles'] = gpon.get('geoFiles') or []
    GPON['vgFile'] = gpon.get('vgFile')


def _touch_routes():
    _save_state()


# ============================ KML PARSER ============================
def _strip_ns(tag: str) -> str:
    return tag.split('}', 1)[1] if '}' in tag else tag


def _parse_kml_bytes(data: bytes):
    result = []
    try:
        root = ET.fromstring(data)
    except ET.ParseError as e:
        print('KML parse error:', e)
        return result
    for pm in root.iter():
        if _strip_ns(pm.tag) != 'Placemark':
            continue
        name = 'route'
        for child in pm:
            if _strip_ns(child.tag) == 'name' and child.text:
                name = child.text.strip()
                break
        for ls in pm.iter():
            if _strip_ns(ls.tag) != 'LineString':
                continue
            coords_el = None
            for child in ls.iter():
                if _strip_ns(child.tag) == 'coordinates':
                    coords_el = child
                    break
            if coords_el is None or not coords_el.text:
                continue
            pts = []
            for token in coords_el.text.split():
                parts = token.split(',')
                if len(parts) >= 2:
                    try:
                        lat = float(parts[1]); lon = float(parts[0])
                        pts.append([lat, lon])
                    except ValueError:
                        continue
            if pts:
                result.append((name, pts))
    return result


def parse_route_file(path, ext):
    if ext == '.kmz':
        try:
            with zipfile.ZipFile(path) as z:
                for n in z.namelist():
                    if n.lower().endswith('.kml'):
                        return _parse_kml_bytes(z.read(n))
        except zipfile.BadZipFile:
            return []
        return []
    with open(path, 'rb') as f:
        return _parse_kml_bytes(f.read())


# ============================ ROUTES ============================
@app.route('/')
def index():
    return render_template('index.html')


@app.route('/upload', methods=['POST'])
def upload():
    global _next_id
    files = request.files.getlist('files[]')
    paths = request.form.getlist('paths[]')
    added = 0
    for i, f in enumerate(files):
        fname = f.filename or f'file_{i}'
        ext = os.path.splitext(fname)[1].lower()
        if ext not in ('.kmz', '.kml'):
            continue
        rel = paths[i] if i < len(paths) else fname
        parts = [p for p in rel.split('/') if p]
        folder = '/'.join(parts[1:-1]) if len(parts) > 1 else ''
        safe = fname.replace('/', '_').replace('\\', '_')
        tmp = os.path.join(UPLOAD_DIR, f'{_next_id}_{safe}')
        f.save(tmp)
        parsed = parse_route_file(tmp, ext)
        multi = len(parsed) > 1
        base = os.path.splitext(fname)[0]
        for name, pts in parsed:
            rid = _next_id
            _next_id += 1
            ROUTES[rid] = {
                'id': rid,
                'name': f'{base} — {name}' if multi else name,
                'branch': None,
                'folder': folder,
                'visible': True,
                'points': pts,
            }
            added += 1
    if added:
        _touch_routes()
    return jsonify({'added': added})


@app.route('/routes')
def routes_list():
    return jsonify(list(ROUTES.values()))


@app.route('/route/<int:rid>', methods=['PATCH'])
def patch_route(rid):
    if rid not in ROUTES:
        return jsonify({'error': 'not found'}), 404
    data = request.get_json(force=True) or {}
    changed = False
    for k in ('name', 'branch', 'folder', 'visible'):
        if k in data:
            ROUTES[rid][k] = data[k]
            changed = True
    if changed:
        _touch_routes()
    return jsonify(ROUTES[rid])


@app.route('/route/<int:rid>', methods=['DELETE'])
def delete_route(rid):
    if ROUTES.pop(rid, None) is not None:
        _touch_routes()
    return jsonify({'ok': True})


@app.route('/merge', methods=['POST'])
def merge():
    global _next_id
    data = request.get_json(force=True) or {}
    ids = data.get('ids', [])
    name = data.get('name', 'Объединённый маршрут')
    pts = []
    for rid in ids:
        r = ROUTES.get(rid)
        if r:
            pts.extend(r['points'])
    if not pts:
        return jsonify({'error': 'no points'}), 400
    rid = _next_id; _next_id += 1
    ROUTES[rid] = {
        'id': rid, 'name': name, 'branch': name,
        'folder': '', 'visible': True, 'points': pts,
    }
    for old in ids:
        ROUTES.pop(old, None)
    _touch_routes()
    return jsonify(ROUTES[rid])


@app.route('/export', methods=['POST'])
def export():
    data = request.get_json(force=True) or {}
    ids = data.get('ids', [])
    group_by_branch = bool(data.get('group_by_branch'))
    kml = simplekml.Kml()
    folders = {}
    for rid in ids:
        r = ROUTES.get(rid)
        if not r:
            continue
        if group_by_branch and r.get('branch'):
            if r['branch'] not in folders:
                folders[r['branch']] = kml.newfolder(name=r['branch'])
            container = folders[r['branch']]
        else:
            container = kml
        ls = container.newlinestring(name=r['name'])
        ls.coords = [(p[1], p[0]) for p in r['points']]
    buf = io.BytesIO()
    kml.savekmz(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.google-earth.kmz',
                     as_attachment=True, download_name='routes.kmz')


# ============================ FOLDERS ============================
@app.route('/folders', methods=['GET'])
def folders_get():
    return jsonify(sorted(PERSIST['customFolders']))


@app.route('/folders', methods=['POST'])
def folders_add():
    data = request.get_json(force=True) or {}
    path = (data.get('path') or '').strip().strip('/')
    if not path:
        return jsonify({'error': 'empty'}), 400
    PERSIST['customFolders'].add(path)
    _touch_routes()
    return jsonify({'ok': True, 'folders': sorted(PERSIST['customFolders'])})


@app.route('/folders', methods=['DELETE'])
def folders_delete():
    data = request.get_json(force=True) or {}
    path = (data.get('path') or '').strip().strip('/')
    if not path:
        return jsonify({'error': 'empty'}), 400
    PERSIST['customFolders'] = {
        f for f in PERSIST['customFolders']
        if f != path and not f.startswith(path + '/')
    }
    _touch_routes()
    return jsonify({'ok': True, 'folders': sorted(PERSIST['customFolders'])})


@app.route('/folders/rename', methods=['POST'])
def folders_rename():
    data = request.get_json(force=True) or {}
    old = (data.get('old') or '').strip().strip('/')
    new = (data.get('new') or '').strip().strip('/')
    if not old or not new:
        return jsonify({'error': 'empty'}), 400
    upd = set()
    for f in PERSIST['customFolders']:
        if f == old:
            upd.add(new)
        elif f.startswith(old + '/'):
            upd.add(new + f[len(old):])
        else:
            upd.add(f)
    PERSIST['customFolders'] = upd
    _touch_routes()
    return jsonify({'ok': True, 'folders': sorted(PERSIST['customFolders'])})


# ============================ STATE EXPORT / IMPORT ============================
@app.route('/state/export', methods=['GET'])
def state_export():
    if not os.path.exists(STATE_FILE):
        return jsonify({'routes': [], 'customFolders': [], 'next_id': 1,
                        'gpon': {'geo': {}, 'onus': [], 'geoFiles': [], 'vgFile': None}})
    return send_file(STATE_FILE, mimetype='application/json',
                     as_attachment=True, download_name='state.json')


@app.route('/state/import', methods=['POST'])
def state_import():
    global _next_id
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'no file'}), 400
    try:
        payload = json.loads(f.read().decode('utf-8'))
    except Exception as e:
        return jsonify({'error': str(e)}), 400
    with _state_lock:
        tmp = STATE_FILE + '.tmp'
        with open(tmp, 'w', encoding='utf-8') as out:
            json.dump(payload, out, ensure_ascii=False)
        os.replace(tmp, STATE_FILE)
    _load_state()
    return jsonify({'ok': True,
                    'routes': len(ROUTES),
                    'folders': len(PERSIST['customFolders']),
                    'geo': len(GPON['geo']),
                    'onus': len(GPON['onus'])})


# ============================ GPON (server-side) ============================
# Само состояние GPON объявлено выше (см. блок STATE), чтобы _save_state()
# и _load_state() могли его сериализовать/восстанавливать.

def _norm_addr(s: str) -> str:
    if not s:
        return ''
    s = str(s).lower().replace('ё', 'е')
    for a, b in [('«', ''), ('»', ''), ('"', ''), ("'", ''), ('`', '')]:
        s = s.replace(a, b)
    while '  ' in s:
        s = s.replace('  ', ' ')
    return s.strip()


def _branch_from_login(login: str) -> str:
    if not login:
        return 'Прочее'
    head = login.split(':')[0]
    parts = head.split('/')
    if len(parts) < 5:
        return head
    return f'{parts[0]}/{parts[3]}/{parts[4]}'


@app.route('/gpon/geo', methods=['POST'])
def gpon_geo():
    if load_workbook is None:
        return jsonify({'error': 'openpyxl not installed'}), 500
    files = request.files.getlist('files[]')
    out = {}
    names = []
    for f in files:
        names.append(f.filename or '')
        try:
            wb = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            rows = list(ws.iter_rows(values_only=True))
            h = -1
            for i, r in enumerate(rows[:30]):
                if r and any(str(c).strip() == 'Широта' for c in r if c is not None):
                    h = i
                    break
            if h < 0:
                continue
            header = [str(c).strip() if c is not None else '' for c in rows[h]]
            cLat = header.index('Широта') if 'Широта' in header else -1
            cLng = header.index('Долгота') if 'Долгота' in header else -1
            cAddr = header.index('Описание') if 'Описание' in header else -1
            if cLat < 0 or cLng < 0 or cAddr < 0:
                continue
            for r in rows[h + 1:]:
                if not r:
                    continue
                try:
                    lat = float(str(r[cLat]).replace(',', '.'))
                    lng = float(str(r[cLng]).replace(',', '.'))
                except (TypeError, ValueError):
                    continue
                addr = r[cAddr]
                if not addr:
                    continue
                out[_norm_addr(addr)] = {'lat': lat, 'lng': lng, 'address': str(addr)}
        except Exception as e:
            print('geo parse error', f.filename, e)
    GPON['geo'] = out
    GPON['geoFiles'] = names
    _save_state()
    return jsonify({'count': len(out), 'files': names})


@app.route('/gpon/vgroups', methods=['POST'])
def gpon_vgroups():
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'no file'}), 400
    raw = f.read()
    text = None
    for enc in ('utf-8-sig', 'utf-8', 'cp1251'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        return jsonify({'error': 'encoding'}), 400

    first = text.split('\n', 1)[0]
    delim = max([';', '\t', ','], key=lambda d: first.count(d))
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = list(reader)
    if not rows:
        return jsonify({'error': 'empty'}), 400

    header = [h.strip() for h in rows[0]]
    i_login = header.index('Логин') if 'Логин' in header else 1
    i_id = header.index('ID') if 'ID' in header else 0
    i_addr = header.index('Адрес') if 'Адрес' in header else 8
    i_stat = next((i for i, h in enumerate(header) if h.startswith('Блокировка')), -1)

    onus = []
    matched = 0
    for r in rows[1:]:
        if not r:
            continue
        try:
            login = r[i_login].strip()
            addr = r[i_addr].strip()
        except IndexError:
            continue
        if not login or not addr or login == 'null' or addr == 'null':
            continue
        g = GPON['geo'].get(_norm_addr(addr))
        status = r[i_stat].strip() if i_stat >= 0 and i_stat < len(r) else ''
        o = {
            'id': r[i_id].strip() if i_id < len(r) else '',
            'login': login,
            'address': addr,
            'status': status,
            'branch': _branch_from_login(login),
            'lat': g['lat'] if g else None,
            'lng': g['lng'] if g else None,
        }
        if g:
            matched += 1
        onus.append(o)
    GPON['onus'] = onus
    GPON['vgFile'] = f.filename
    _save_state()
    return jsonify({'total': len(onus), 'matched': matched,
                    'unmatched': len(onus) - matched})


# ============================ GPON STATE ============================
@app.route('/gpon/state', methods=['GET'])
def gpon_state_get():
    with _state_lock:
        return jsonify({
            'geo': GPON['geo'],
            'onus': GPON['onus'],
            'geoFiles': GPON.get('geoFiles', []),
            'vgFile': GPON.get('vgFile'),
        })


@app.route('/gpon/state', methods=['POST'])
def gpon_state_post():
    data = request.get_json(force=True, silent=True) or {}
    with _state_lock:
        if 'geo' in data and isinstance(data['geo'], dict):
            GPON['geo'] = data['geo']
        if 'onus' in data and isinstance(data['onus'], list):
            GPON['onus'] = data['onus']
        if 'geoFiles' in data:
            GPON['geoFiles'] = list(data.get('geoFiles') or [])
        if 'vgFile' in data:
            GPON['vgFile'] = data.get('vgFile')
    _save_state()
    return jsonify({
        'ok': True,
        'geo': len(GPON['geo']),
        'onus': len(GPON['onus']),
    })


@app.route('/gpon/state', methods=['DELETE'])
def gpon_state_delete():
    with _state_lock:
        GPON['geo'] = {}
        GPON['onus'] = []
        GPON['geoFiles'] = []
        GPON['vgFile'] = None
    _save_state()
    return jsonify({'ok': True})


# ============================ BOOT ============================
_load_state()

if __name__ == '__main__':
    app.run(debug=True, host='127.0.0.1', port=5000)